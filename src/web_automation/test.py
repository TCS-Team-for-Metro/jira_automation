from __future__ import annotations

from base64 import b64encode
import json
import sys
from csv import reader
from datetime import datetime
from pathlib import Path
import re
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen
from uuid import uuid4

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from playwright.sync_api import Error as PlaywrightError
from playwright.sync_api import sync_playwright
try:
    from .env_config import get_required_env
except ImportError:
    from env_config import get_required_env

JIRA_ISSUES_URL = "https://jira.metro.digital/login.jsp"
CSV_DOWNLOAD_URL = "https://jira.metro.digital/sr/jira.issueviews:searchrequest-csv-current-fields/139312/SearchRequest-139312.csv"
MICROSOFT_EMAIL = "ashish.dake@metro-external.digital"

TEAMS_WEBHOOK_URL = get_required_env("TEAMS_WEBHOOK_TEST")

COLUMNS_TO_REMOVE = {
    "Issue id",
    "Project key",
    "Project type",
    "Project lead",
    "Project description",
    "Project url",
}
ISSUE_KEY_COLUMN_INDEX = 0
UNASSIGNED_VALUE = "Unassigned"
WORK_IN_PROGRESS_STATUS_PARTS = ("open", "delegated", "feedback received", "in analysis")
PENDING_STATUS = "pending"
WAITING_FOR_USER_STATUS = "waiting for user"
WAITING_FOR_ACCEPTANCE_STATUS = "waiting for acceptance"
CLOSED_STATUS = "closed"
HEADER_RENAMES = {
    "Custom field (Product)": "Product",
    "Custom field (Solution)": "Solution",
}
PROJECT_NAME_HEADER = "Project name"
ISSUE_KEY_HEADER = "Issue key"
PRODUCT_HEADER = "Product"
SOLUTION_HEADER = "Solution"
LABELS_HEADER = "Labels"
PREFERRED_COLUMN_ORDER = [
    "Project name",
    "Solution",
    "Product",
    "Issue key",
    "Summary",
    "Assignee",
    "Reporter",
    "Priority",
    "Status",
    "Created",
    "Updated",
    "Time to resolution",
    "Labels",
]
USER_STATUS_REPORT_USERS = [
    "Ashish Dake",
    "Rajkumar Kale",
    "Narayan Mhaske",
    "Pranav Soan",
    "Rashmi Dubey",
    "Deepak Gavel",
    "Sajida Mullani",
    "Mahesh Pujari",
    "Ramamurthy Rongala",
    "Namrata Shandilya",
    "Rishabh Verma",
    "Pooja Malage",
]

TARGET_TICKET_KEYS = [
    "SDMCCCS-131025",
    "SDMCCCS-130958",
    "SDMCCCS-130900",
    "SDMCCCS-130851",
    "SDMCCCS-130752",
    "SDMCCCS-130585",
    "SDMCCCS-130071",
    "SDMCCCS-129968",
    "SDMCCCS-129985",
    "SDMCCCS-130022",
    "SDMCCCS-129937",
    "SDMCCCS-129809",
    "SDMCCCS-129077",
    "SDMCCCS-129560",
    "SDMCCCS-130366",
    "SDMCCCS-129946",
    "SDMCCCS-129927",
    "SDMCCCS-129696",
    "SDMCCCS-129131",
    "SDMCCCS-129019",
    "SDMCCCS-129462",
    "SDMCCCS-130385",
    "SDMCCCS-130951",
    "SDMCCCS-130809",
    "SDMCCCS-130614",
    "SDMCCCS-130619",
    "SDMCCCS-130938",
    "SDMCCCS-130488",
    "SDMCCCS-130476",
    "SDMCCCS-130317",
    "SDMCCCS-131059",
    "SDMCCCS-131393",
    "SDMCCCS-131366",
    "SDMCCCS-131210",
    "SDMCCCS-131208",
    "SDMCCCS-131516",
    "SDMCCCS-131623",
    "SDMCCCS-131724",
    "SDMCCCS-131723",
    "SDMCCCS-131722",
    "SDMCCCS-131773",
    "SDMCCCS-131774",
    "SDMCCCS-131973",
    "SDMCCCS-131963",
    "SDMCCCS-131965",
    "SDMCCCS-132091",
    "SDMCCCS-132149",
]

TICKET_KEY_ORDER = {
    ticket_key: index for index, ticket_key in enumerate(TARGET_TICKET_KEYS)
}
TARGET_TICKET_KEY_SET = set(TARGET_TICKET_KEYS)


def build_download_filename() -> str:
    print("Building CSV download filename.")
    today = datetime.now().strftime("%d-%b-%Y")
    return f"Wave 3 Daily Tickets {today}.csv"


def build_output_filename() -> str:
    print("Building XLSX output filename.")
    today = datetime.now().strftime("%d-%b-%Y")
    return f"Wave 3 Daily Tickets {today}.xlsx"


def is_jira_url(url: str) -> bool:
    print(f"Checking whether URL is Jira: {url}")
    return "jira.metro.digital" in url


def is_microsoft_login_url(url: str) -> bool:
    print(f"Checking whether URL is a Microsoft login page: {url}")
    return "login.microsoftonline.com" in url or "login.live.com" in url


def build_columns_to_remove_indexes(header_row: list[str]) -> set[int]:
    print("Building indexes for unwanted CSV columns.")
    normalized_columns_to_remove = {
        column_name.strip().lower() for column_name in COLUMNS_TO_REMOVE
    }
    return {
        index
        for index, column_name in enumerate(header_row)
        if column_name.strip().lower() in normalized_columns_to_remove
    }


def remove_columns(row: list[str], columns_to_remove_indexes: set[int]) -> list[str]:
    print("Removing unwanted CSV columns from row.")
    return [
        value
        for index, value in enumerate(row)
        if index not in columns_to_remove_indexes
    ]


def filter_row(row: list[str], columns_to_remove_indexes: set[int]) -> list[str]:
    print("Filtering unwanted CSV columns from row.")
    return remove_columns(row, columns_to_remove_indexes)


def rename_columns(row: list[str]) -> list[str]:
    print("Renaming CSV column names.")
    return [HEADER_RENAMES.get(value.strip(), value) for value in row]


def rename_headers(row: list[str]) -> list[str]:
    print("Renaming CSV headers.")
    return rename_columns(row)


def move_project_name_to_start(row: list[str], project_name_index: int | None) -> list[str]:
    print("Moving Project name column to the start of the row.")
    if project_name_index is None or project_name_index >= len(row):
        return row

    reordered_row = row.copy()
    reordered_row.insert(0, reordered_row.pop(project_name_index))
    return reordered_row


def swap_product_and_solution(row: list[str], product_index: int | None, solution_index: int | None) -> list[str]:
    print("Swapping Product and Solution columns.")
    if (
        product_index is None
        or solution_index is None
        or product_index >= len(row)
        or solution_index >= len(row)
    ):
        return row

    reordered_row = row.copy()
    reordered_row[product_index], reordered_row[solution_index] = (
        reordered_row[solution_index],
        reordered_row[product_index],
    )
    return reordered_row


def reposition_columns(
    row: list[str],
    project_name_index: int | None,
    product_index: int | None,
    solution_index: int | None,
) -> list[str]:
    print("Re-positioning columns in the row.")
    reordered_row = move_project_name_to_start(row, project_name_index)
    return swap_product_and_solution(reordered_row, product_index, solution_index)


def build_column_indexes_by_name(row: list[str]) -> dict[str, list[int]]:
    indexes_by_name: dict[str, list[int]] = {}
    for index, column_name in enumerate(row):
        normalized_name = column_name.strip().lower()
        indexes_by_name.setdefault(normalized_name, []).append(index)
    return indexes_by_name


def build_preferred_column_indexes(
    header_row: list[str], preferred_column_order: list[str]
) -> list[int]:
    print("Building preferred column indexes from header row.")
    indexes_by_name = build_column_indexes_by_name(header_row)

    preferred_indexes: list[int] = []
    for preferred_name in preferred_column_order:
        normalized_name = preferred_name.strip().lower()
        matching_indexes = indexes_by_name.get(normalized_name)
        if not matching_indexes:
            continue
        preferred_indexes.append(matching_indexes.pop(0))

    return preferred_indexes


def reorder_row_by_indexes(row: list[str], column_indexes: list[int]) -> list[str]:
    print("Arranging row values in preferred column order.")
    reordered_row: list[str] = []
    for index in column_indexes:
        if index < len(row):
            reordered_row.append(row[index])
        else:
            reordered_row.append("")
    return reordered_row


def merge_columns_by_comma(row: list[str], source_indexes: list[int]) -> str:
    values: list[str] = []
    for index in source_indexes:
        if index >= len(row):
            continue
        value = row[index].strip()
        if value:
            values.append(value)
    return ", ".join(values)


def set_default_assignee(row: list[str], assignee_index: int | None) -> list[str]:
    print("Normalizing Assignee value.")
    if assignee_index is None or assignee_index >= len(row):
        return row

    if not row[assignee_index].strip():
        row[assignee_index] = UNASSIGNED_VALUE
    return row


def has_assigned_owner(row: list[str], assignee_index: int | None) -> bool:
    print("Checking whether issue has an assigned owner.")
    if assignee_index is None or assignee_index >= len(row):
        return False

    assignee = row[assignee_index].strip()
    return bool(assignee) and assignee.lower() != UNASSIGNED_VALUE.lower()


def create_status_counts() -> dict[str, int]:
    print("Creating status counter.")
    return {
        "work_in_progress": 0,
        "pending": 0,
        "waiting_for_user": 0,
        "waiting_for_acceptance": 0,
        "closed": 0,
    }


def create_user_status_counts() -> dict[str, dict[str, int]]:
    return {user: create_status_counts() for user in USER_STATUS_REPORT_USERS}


def calculate_eod_ticket_count(status_counts: dict[str, int]) -> int:
    return (
        status_counts["work_in_progress"]
        + status_counts["pending"]
        + status_counts["waiting_for_user"]
        + status_counts["waiting_for_acceptance"]
    )


def normalize_letters(value: str) -> str:
    return re.sub(r"[^a-z]", "", value.lower())


def get_user_name_parts(user_name: str) -> tuple[str, str]:
    parts = re.findall(r"[a-z]+", user_name.lower())
    if not parts:
        return "", ""
    return parts[0], parts[-1]


def resolve_report_user_from_assignee(assignee: str) -> str | None:
    normalized_assignee = normalize_letters(assignee)
    if not normalized_assignee:
        return None

    for user_name in USER_STATUS_REPORT_USERS:
        first_name, last_name = get_user_name_parts(user_name)
        if first_name and last_name and first_name in normalized_assignee and last_name in normalized_assignee:
            return user_name

    return None


def update_user_status_counts(
    row: list[str],
    assignee_index: int | None,
    status_index: int | None,
    user_status_counts: dict[str, dict[str, int]],
) -> None:
    if (
        assignee_index is None
        or assignee_index >= len(row)
        or status_index is None
        or status_index >= len(row)
    ):
        return

    assignee = row[assignee_index].strip()
    if not assignee:
        return

    report_user = resolve_report_user_from_assignee(assignee)
    if report_user is None:
        return

    update_status_counts(row, status_index, user_status_counts[report_user])


def update_status_counts(
    row: list[str], status_index: int | None, status_counts: dict[str, int]
) -> None:
    print("Updating status counters.")
    if status_index is None or status_index >= len(row):
        return

    status = row[status_index].strip().lower()
    if any(part in status for part in WORK_IN_PROGRESS_STATUS_PARTS):
        status_counts["work_in_progress"] += 1
    elif status == PENDING_STATUS:
        status_counts["pending"] += 1
    elif status == WAITING_FOR_USER_STATUS:
        status_counts["waiting_for_user"] += 1
    elif status == WAITING_FOR_ACCEPTANCE_STATUS:
        status_counts["waiting_for_acceptance"] += 1
    elif status == CLOSED_STATUS:
        status_counts["closed"] += 1


def get_known_issue_prefixes() -> tuple[str, ...]:
    return ("SDMCCCS", "SDMCCUA")


def is_effectively_empty_row(row: list[str]) -> bool:
    return not row or all((value is None or str(value).strip() == "") for value in row)


def print_status_counts(sheet_name: str, status_counts: dict[str, int]) -> None:
    print(f"{sheet_name} counts:")
    print(f"  Work in progress: {status_counts['work_in_progress']}")
    print(f"  Pending: {status_counts['pending']}")
    print(f"  Waiting for user: {status_counts['waiting_for_user']}")
    print(f"  Waiting for acceptance: {status_counts['waiting_for_acceptance']}")
    print(f"  Closed: {status_counts['closed']}")


def print_user_status_counts_table(user_status_counts: dict[str, dict[str, int]]) -> None:
    headers = [
        "User",
        "Work In Progress",
        "Waiting For User",
        "Pending",
        "Waiting For Acceptance",
        "Closed",
        "EoD Ticket Count",
    ]
    rows = []
    for user_name in USER_STATUS_REPORT_USERS:
        status_counts = user_status_counts[user_name]
        rows.append(
            [
                user_name,
                str(status_counts["work_in_progress"]),
                str(status_counts["waiting_for_user"]),
                str(status_counts["pending"]),
                str(status_counts["waiting_for_acceptance"]),
                str(status_counts["closed"]),
                str(calculate_eod_ticket_count(status_counts)),
            ]
        )

    widths = [len(header) for header in headers]
    for row in rows:
        for index, cell_value in enumerate(row):
            widths[index] = max(widths[index], len(cell_value))

    separator = "+-" + "-+-".join("-" * width for width in widths) + "-+"

    def format_row(cells: list[str]) -> str:
        return "| " + " | ".join(
            cell_value.ljust(widths[index]) for index, cell_value in enumerate(cells)
        ) + " |"

    print("User status table:")
    print(separator)
    print(format_row(headers))
    print(separator)
    for row in rows:
        print(format_row(row))
    print(separator)


def send_status_counts_to_teams_webhook(
    czsk_status_counts: dict[str, int],
    ukraine_status_counts: dict[str, int],
    user_status_counts: dict[str, dict[str, int]],
    xlsx_path: Path | None = None,
) -> None:
    print("Sending status counts to Teams webhook.")
    users_payload: dict[str, dict[str, int]] = {}
    for user_name, status_counts in user_status_counts.items():
        users_payload[user_name] = {
            **status_counts,
            "eod_ticket_count": calculate_eod_ticket_count(status_counts),
        }

    payload = {
        "timestamp": datetime.now().isoformat(timespec="seconds"),
        "czsk": czsk_status_counts,
        "ukraine": ukraine_status_counts,
        "users": users_payload,
    }
    if xlsx_path is not None:
        payload["generated_xlsx"] = {
            "file_name": xlsx_path.name,
            "mime_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "content_base64": b64encode(xlsx_path.read_bytes()).decode("ascii"),
        }

    request = Request(
        TEAMS_WEBHOOK_URL,
        data=json.dumps(payload).encode("utf-8"),
        headers={"Content-Type": "application/json"},
        method="POST",
    )
    try:
        with urlopen(request, timeout=30):
            pass
    except HTTPError as exc:
        raise RuntimeError(
            f"Teams webhook returned HTTP {exc.code}: {exc.reason}"
        ) from exc
    except URLError as exc:
        raise RuntimeError(f"Teams webhook request failed: {exc.reason}") from exc


def format_worksheet(worksheet) -> None:
    print(f"Formatting worksheet: {worksheet.title}")
    header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    header_font = Font(bold=True)

    if worksheet.max_row == 0:
        return

    for cell in worksheet[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    worksheet.column_dimensions["C"].width = 40

    for column_index in range(1, worksheet.max_column + 1):
        column_letter = get_column_letter(column_index)
        if column_letter == "C":
            continue

        max_length = 0
        for row_index in range(1, worksheet.max_row + 1):
            value = worksheet.cell(row=row_index, column=column_index).value
            if value is None:
                continue
            max_length = max(max_length, len(str(value)))

        worksheet.column_dimensions[column_letter].width = max_length + 2


def convert_csv_to_xlsx(csv_path: Path) -> Path:
    print(f"Converting CSV to XLSX: {csv_path}")
    workbook = Workbook()
    tickets_sheet = workbook.active
    tickets_sheet.title = "Tickets"
    issue_key_index: int | None = None
    rows_read = 0
    data_rows_written = 0
    matched_rows: list[list[str]] = []
    found_ticket_keys: set[str] = set()

    with csv_path.open("r", encoding="utf-8-sig", newline="") as csv_file:
        for row_number, row in enumerate(reader(csv_file), start=1):
            rows_read += 1
            if is_effectively_empty_row(row):
                print(f"Skipping empty row {row_number}.")
                continue

            if row_number == 1:
                issue_key_index = next(
                    (
                        index
                        for index, column_name in enumerate(row)
                        if column_name.strip().lower() == ISSUE_KEY_HEADER.lower()
                    ),
                    None,
                )
                if issue_key_index is None:
                    raise RuntimeError("Issue key column not found in downloaded CSV header.")
                tickets_sheet.append(row)
                continue

            if issue_key_index is None or issue_key_index >= len(row):
                print(
                    f"Skipping row {row_number} because Issue key column is not available in transformed row."
                )
                continue

            issue_key = row[issue_key_index].strip().upper()
            if not issue_key:
                print(f"Skipping row {row_number} because Issue key is empty.")
                continue

            if issue_key in TARGET_TICKET_KEY_SET:
                matched_rows.append(row)
                found_ticket_keys.add(issue_key)
                data_rows_written += 1

    print(f"CSV rows read: {rows_read}")
    print(f"Matched ticket rows: {data_rows_written}")

    matched_rows.sort(
        key=lambda row: TICKET_KEY_ORDER.get(row[issue_key_index].strip().upper(), len(TARGET_TICKET_KEYS))
    )
    for matched_row in matched_rows:
        tickets_sheet.append(matched_row)

    missing_ticket_keys = [
        ticket_key for ticket_key in TARGET_TICKET_KEYS if ticket_key not in found_ticket_keys
    ]
    if missing_ticket_keys:
        print(f"Tickets not found in CSV ({len(missing_ticket_keys)}): {', '.join(missing_ticket_keys)}")

    xlsx_path = csv_path.with_name(build_output_filename())
    print(f"Saving XLSX file: {xlsx_path}")
    workbook.save(xlsx_path)
    print(f"Removing original CSV file: {csv_path}")
    csv_path.unlink()
    return xlsx_path


def fill_microsoft_email(page) -> bool:
    print("Attempting to fill Microsoft email field.")
    selectors = [
        "input[type='email'][name='loginfmt']",
        "input[name='loginfmt']",
        "input[type='email']",
        "input#i0116",
        "input[type='text'][name='loginfmt']",
        "input[type='text']",
    ]

    for selector in selectors:
        locator = page.locator(selector)
        if locator.count() == 0:
            continue

        field = locator.first
        if not field.is_visible():
            continue

        field.click()
        field.press("Control+A")
        field.fill(MICROSOFT_EMAIL)
        print(f"Microsoft email field filled using selector: {selector}")
        return True

    print("Microsoft email field is not visible.")
    return False


def click_primary_button(page, *texts: str):
    print(f"Trying primary buttons: {texts}")
    for text in texts:
        locator = page.locator(
            f"button:has-text('{text}'), input[type='submit'][value='{text}'], "
            f"[type='submit'][value='{text}'], [role='button']:has-text('{text}')"
        )
        if locator.count() > 0:
            button = locator.first
            if button.is_visible():
                button.click()
                print(f"Clicked primary button: {text}")
                return True
    print("No matching primary button found.")
    return False


def login_to_microsoft(page) -> None:
    print("Starting Jira login flow.")
    page.goto(JIRA_ISSUES_URL, wait_until="domcontentloaded", timeout=120000)

    deadline = datetime.now().timestamp() + 120
    while datetime.now().timestamp() < deadline:
        current_url = page.url.lower()
        if is_jira_url(current_url):
            print("Jira page detected. Login flow complete.")
            return

        if is_microsoft_login_url(current_url):
            if fill_microsoft_email(page):
                if not click_primary_button(page, "Next", "Continue"):
                    print("Falling back to default submit button.")
                    submit_buttons = page.locator("input[type='submit'], button[type='submit']")
                    if submit_buttons.count() > 0 and submit_buttons.first.is_visible():
                        submit_buttons.first.click()
                    else:
                        page.locator("input[type='submit']").first.click()
                page.wait_for_timeout(2000)
                continue

            print("Waiting for Microsoft login page to finish loading.")
            page.wait_for_load_state("domcontentloaded", timeout=15000)
            page.wait_for_timeout(1000)
            continue

        page.wait_for_timeout(1000)

    raise RuntimeError("Microsoft login flow did not reach the Jira dashboard.")


def download_csv(page) -> Path:
    print("Starting CSV download.")
    download_dir = Path.home() / "Downloads"
    download_dir.mkdir(parents=True, exist_ok=True)

    file_name = build_download_filename()
    temp_file = download_dir / f"{uuid4()}-{file_name}"

    with page.expect_download() as download_info:
        page.evaluate("url => window.location.href = url", CSV_DOWNLOAD_URL)

    download = download_info.value
    download.save_as(path=temp_file)

    final_path = download_dir / file_name
    temp_file.replace(final_path)
    print(f"CSV downloaded to: {final_path}")
    return final_path


def main() -> int:
    print("Launching browser automation.")
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=False)
        context = browser.new_context(accept_downloads=True)
        page = context.new_page()

        try:
            login_to_microsoft(page)
            csv_path = download_csv(page)
            final_path = convert_csv_to_xlsx(csv_path)
            print(f"Downloaded XLSX to: {final_path}")
            return 0
        except PlaywrightError as exc:
            print(f"Playwright error: {exc}", file=sys.stderr)
            return 1
        except Exception as exc:  # pragma: no cover
            print(f"Failed to download Jira CSV: {exc}", file=sys.stderr)
            return 1
        finally:
            print("Closing browser.")
            browser.close()


if __name__ == "__main__":
    raise SystemExit(main())
